
"""Load the Term/Metric workbook sheets and insert hybrid Qdrant Points."""

from __future__ import annotations

import argparse
import json
import os
from collections.abc import Callable, Iterator, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import NAMESPACE_URL, uuid5

from dotenv import load_dotenv
from openpyxl import load_workbook
from qdrant_client import QdrantClient, models

from vector_db.text2vector import HybridEmbedding, texts_to_hybrid_vectors


DEFAULT_WORKBOOK_PATH = Path("DOCS/knowledge_base.xlsx")
COLLECTION_NAME = "knowledge_base"
DENSE_VECTOR_NAME = "knowledge_dense"
SPARSE_VECTOR_NAME = "knowledge_sparse"
VECTOR_DIMENSION = 1024
DEFAULT_EMBEDDING_BUFFER_SIZE = 512
DEFAULT_UPLOAD_BATCH_SIZE = 128
SHEET_KNOWLEDGE_TYPES = {"Term": "TERM", "Metric": "METRIC"}
REQUIRED_COLUMNS = (
    "knowledge_type",
    "canonical_term",
    "aliases",
    "description",
    "dataset",
    "category",
)
POINT_ID_NAMESPACE = uuid5(NAMESPACE_URL, "miraeasset:qdrant:knowledge_base")

BatchVectorizer = Callable[[Sequence[str]], list[HybridEmbedding]]


@dataclass(frozen=True, slots=True)
class KnowledgeItem:
    knowledge_type: str
    name: str
    aliases: tuple[str, ...]
    description: str
    datasets: tuple[str, ...]
    categories: tuple[str, ...]

    def payload(self) -> dict[str, Any]:
        return {
            "knowledge_type": self.knowledge_type,
            "name": self.name,
            "aliases": list(self.aliases),
            "description": self.description,
            "datasets": list(self.datasets),
            "categories": list(self.categories),
        }


@dataclass(frozen=True, slots=True)
class KnowledgePointInput:
    id: str
    embedding_text: str
    payload: dict[str, Any]


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _split_values(value: Any) -> tuple[str, ...]:
    values: list[str] = []
    seen: set[str] = set()
    for part in _text(value).split("|"):
        normalized = part.strip()
        if normalized and normalized not in seen:
            values.append(normalized)
            seen.add(normalized)
    return tuple(values)


def _required_text(
    row: Mapping[str, Any],
    column: str,
    *,
    sheet_name: str,
    row_number: int,
) -> str:
    value = _text(row.get(column))
    if not value:
        raise ValueError(
            f"{sheet_name}!{row_number} requires a non-empty {column} value"
        )
    return value


def _sheet_items(
    sheet: Any,
    *,
    expected_knowledge_type: str,
) -> Iterator[KnowledgeItem]:
    rows = sheet.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration as error:
        raise ValueError(f"Workbook sheet is empty: {sheet.title}") from error

    headers = [_text(value) for value in header_values]
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing_columns:
        raise ValueError(
            f"{sheet.title} is missing required columns: "
            + ", ".join(missing_columns)
        )

    for row_number, values in enumerate(rows, start=2):
        row = dict(zip(headers, values, strict=False))
        if not any(_text(value) for value in values):
            continue
        knowledge_type = _required_text(
            row,
            "knowledge_type",
            sheet_name=sheet.title,
            row_number=row_number,
        ).upper()
        if knowledge_type != expected_knowledge_type:
            raise ValueError(
                f"{sheet.title}!{row_number} has unexpected knowledge_type: "
                f"{knowledge_type}"
            )
        aliases = _split_values(row.get("aliases"))
        datasets = _split_values(row.get("dataset"))
        categories = _split_values(row.get("category"))
        if not datasets:
            raise ValueError(f"{sheet.title}!{row_number} requires a dataset")
        if not categories:
            raise ValueError(f"{sheet.title}!{row_number} requires a category")
        yield KnowledgeItem(
            knowledge_type=knowledge_type,
            name=_required_text(
                row,
                "canonical_term",
                sheet_name=sheet.title,
                row_number=row_number,
            ),
            aliases=aliases,
            description=_required_text(
                row,
                "description",
                sheet_name=sheet.title,
                row_number=row_number,
            ),
            datasets=datasets,
            categories=categories,
        )


def load_knowledge_items(path: Path = DEFAULT_WORKBOOK_PATH) -> list[KnowledgeItem]:
    """Read Term and Metric rows and remove exact normalized duplicates."""
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Knowledge workbook does not exist: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        missing_sheets = [
            sheet_name
            for sheet_name in SHEET_KNOWLEDGE_TYPES
            if sheet_name not in workbook.sheetnames
        ]
        if missing_sheets:
            raise ValueError(
                "Knowledge workbook is missing sheets: " + ", ".join(missing_sheets)
            )

        unique_items: dict[KnowledgeItem, None] = {}
        for sheet_name, expected_type in SHEET_KNOWLEDGE_TYPES.items():
            for item in _sheet_items(
                workbook[sheet_name],
                expected_knowledge_type=expected_type,
            ):
                unique_items.setdefault(item, None)
        return list(unique_items)
    finally:
        workbook.close()


def build_embedding_text(item: KnowledgeItem) -> str:
    lines = [
        f"지식 유형 : {item.knowledge_type}",
        f"이름 : {item.name}",
    ]
    if item.aliases:
        lines.append(f"별칭 : {' | '.join(item.aliases)}")
    lines.extend(
        [
            f"설명 : {item.description}",
            f"데이터셋 : {' | '.join(item.datasets)}",
            f"카테고리 : {' | '.join(item.categories)}",
        ]
    )
    return "\n".join(lines)


def _point_id(payload: Mapping[str, Any]) -> str:
    identity = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return str(uuid5(POINT_ID_NAMESPACE, identity))


def build_point_inputs(items: Sequence[KnowledgeItem]) -> list[KnowledgePointInput]:
    point_inputs: list[KnowledgePointInput] = []
    for item in items:
        if not isinstance(item, KnowledgeItem):
            raise ValueError("items must contain KnowledgeItem values")
        payload = item.payload()
        point_inputs.append(
            KnowledgePointInput(
                id=_point_id(payload),
                embedding_text=build_embedding_text(item),
                payload=payload,
            )
        )
    return point_inputs


def _chunks(values: Sequence[Any], size: int) -> Iterator[Sequence[Any]]:
    if not isinstance(size, int) or isinstance(size, bool) or size < 1:
        raise ValueError("batch size must be a positive integer")
    for start in range(0, len(values), size):
        yield values[start : start + size]


def _qdrant_client_from_env() -> QdrantClient:
    load_dotenv()
    host = os.getenv("QDRANT_HOST", "").strip()
    port_text = os.getenv("QDRANT_PORT", "").strip()
    if not host or not port_text:
        raise RuntimeError("QDRANT_HOST and QDRANT_PORT must be configured")
    try:
        port = int(port_text)
    except ValueError as error:
        raise RuntimeError("QDRANT_PORT must be an integer") from error
    return QdrantClient(host=host, port=port)


def validate_collection(client: QdrantClient) -> None:
    if not client.collection_exists(COLLECTION_NAME):
        raise RuntimeError(f"Qdrant collection does not exist: {COLLECTION_NAME}")
    collection = client.get_collection(COLLECTION_NAME)
    vectors = collection.config.params.vectors
    dense = vectors.get(DENSE_VECTOR_NAME) if isinstance(vectors, Mapping) else None
    if dense is None or dense.size != VECTOR_DIMENSION:
        raise RuntimeError(
            f"{COLLECTION_NAME}.{DENSE_VECTOR_NAME} must have dimension "
            f"{VECTOR_DIMENSION}"
        )
    if dense.distance != models.Distance.COSINE:
        raise RuntimeError(
            f"{COLLECTION_NAME}.{DENSE_VECTOR_NAME} must use Cosine distance"
        )
    sparse_vectors = collection.config.params.sparse_vectors
    if not isinstance(sparse_vectors, Mapping) or SPARSE_VECTOR_NAME not in sparse_vectors:
        raise RuntimeError(
            f"Qdrant sparse vector is missing: {SPARSE_VECTOR_NAME}"
        )


def _assemble_points(
    point_inputs: Sequence[KnowledgePointInput],
    embeddings: Sequence[HybridEmbedding],
) -> list[models.PointStruct]:
    if len(point_inputs) != len(embeddings):
        raise ValueError("Embedding count does not match Point input count")
    return [
        models.PointStruct(
            id=point_input.id,
            vector={
                DENSE_VECTOR_NAME: list(embedding.dense),
                SPARSE_VECTOR_NAME: models.SparseVector(
                    indices=list(embedding.sparse.indices),
                    values=list(embedding.sparse.values),
                ),
            },
            payload=point_input.payload,
        )
        for point_input, embedding in zip(point_inputs, embeddings, strict=True)
    ]


def insert_point_inputs(
    client: QdrantClient,
    point_inputs: Sequence[KnowledgePointInput],
    *,
    embedding_buffer_size: int = DEFAULT_EMBEDDING_BUFFER_SIZE,
    upload_batch_size: int = DEFAULT_UPLOAD_BATCH_SIZE,
    vectorizer: BatchVectorizer = texts_to_hybrid_vectors,
) -> int:
    inserted = 0
    for embedding_batch in _chunks(point_inputs, embedding_buffer_size):
        embeddings = vectorizer(
            [point_input.embedding_text for point_input in embedding_batch]
        )
        points = _assemble_points(embedding_batch, embeddings)
        for point_batch in _chunks(points, upload_batch_size):
            client.upsert(
                collection_name=COLLECTION_NAME,
                wait=True,
                points=list(point_batch),
            )
            inserted += len(point_batch)
    return inserted


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Insert Term and Metric knowledge into Qdrant."
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK_PATH)
    parser.add_argument(
        "--embedding-buffer-size",
        type=int,
        default=DEFAULT_EMBEDDING_BUFFER_SIZE,
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=DEFAULT_UPLOAD_BATCH_SIZE,
    )
    parser.add_argument("--dry-run", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    items = load_knowledge_items(args.workbook)
    point_inputs = build_point_inputs(items)
    if args.dry_run:
        result = {
            "dry_run": True,
            "workbook": str(args.workbook),
            "points": len(point_inputs),
            "point_kinds": {
                knowledge_type: sum(
                    point.payload["knowledge_type"] == knowledge_type
                    for point in point_inputs
                )
                for knowledge_type in ("TERM", "METRIC")
            },
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0

    client = _qdrant_client_from_env()
    try:
        validate_collection(client)
        inserted = insert_point_inputs(
            client,
            point_inputs,
            embedding_buffer_size=args.embedding_buffer_size,
            upload_batch_size=args.batch_size,
        )
    finally:
        client.close()
    print(
        json.dumps(
            {
                "dry_run": False,
                "workbook": str(args.workbook),
                "points": len(point_inputs),
                "inserted": inserted,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
